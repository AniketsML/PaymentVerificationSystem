"""
31-column audit-grade Excel workbook generator for SARFAESI extraction.
"""
from __future__ import annotations
import io
from typing import Any, Dict, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from workspaces.legal.models import SARFAESILeadResult

_HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
_WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
_CURRENCY_FMT = '₹ #,##0.00'
_CURRENCY_COLS = SARFAESILeadResult.CURRENCY_COLUMNS


def generate_workbook_bytes(leads: List[Dict[str, Any]]) -> bytes:
    """Generate styled Excel workbook and return as bytes."""
    wb = _build_workbook(leads)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_workbook(leads: List[Dict[str, Any]], output_path: str) -> str:
    """Generate styled Excel workbook and save to path."""
    wb = _build_workbook(leads)
    wb.save(output_path)
    return output_path


def _build_workbook(leads: List[Dict[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "SARFAESI Extraction"
    headers = SARFAESILeadResult.EXCEL_COLUMNS
    fields = SARFAESILeadResult.EXCEL_FIELD_MAP

    # Write header row
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER

    # Write data rows
    for row_idx, lead in enumerate(leads, 2):
        fill = _ALT_FILL if row_idx % 2 == 0 else _WHITE_FILL
        for col_idx, fld in enumerate(fields, 1):
            val = lead.get(fld)
            if val is not None and col_idx - 1 in _CURRENCY_COLS:
                try:
                    val = float(str(val).replace(",", ""))
                except (ValueError, TypeError):
                    pass
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
            cell.fill = fill
            cell.border = _BORDER
            if col_idx - 1 in _CURRENCY_COLS:
                cell.number_format = _CURRENCY_FMT
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Auto-fit column widths
    for col_idx in range(1, len(headers) + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row in ws.iter_rows(min_row=2, max_row=min(len(leads) + 1, 50), min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 40))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

    ws.freeze_panes = "A2"
    return wb
